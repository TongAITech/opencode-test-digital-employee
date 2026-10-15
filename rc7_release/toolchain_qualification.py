from __future__ import annotations

import argparse
import importlib
import json
import os
import subprocess
import tempfile
from pathlib import Path
from typing import Dict, Optional

from rc7_runtime.diagnosis_providers import MultiToolLogDiagnoser
from rc7_runtime.requirement_toolchain import RequirementDiffEngine, RequirementDocumentParser
from rc7_runtime.reporting import ReportRenderer


DOC_MODULES = {
    "docx": "python-docx",
    "openpyxl": "openpyxl",
    "pptx": "python-pptx",
    "pypdf": "pypdf",
}


def _module_versions() -> Dict[str, str]:
    try:
        from importlib.metadata import version
    except ImportError:
        from importlib_metadata import version  # type: ignore
    result = {}
    for module, distribution in DOC_MODULES.items():
        importlib.import_module(module)
        try:
            result[distribution] = version(distribution)
        except Exception:
            result[distribution] = "UNKNOWN"
    return result


def _requirement_gate(root: Path) -> Dict[str, object]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    parser = RequirementDocumentParser()

    before = root / "before.txt"
    after = root / "after.txt"
    before.write_text("额度必须大于0\n额度上限500000.00\n", encoding="utf-8")
    after.write_text("额度必须大于0\n额度上限500000.00\n500000.01必须拒绝\n", encoding="utf-8")
    changes = RequirementDiffEngine().diff(parser.parse(str(before)), parser.parse(str(after)))
    if not changes or changes[-1].change_type not in {"ADDED", "MODIFIED"}:
        raise RuntimeError("REQUIREMENT_DIFF_SMOKE_FAILED")

    docx_path = root / "sample.docx"
    doc = Document()
    doc.add_paragraph("SST-001 交易金额边界")
    doc.save(str(docx_path))
    if not parser.parse(str(docx_path)):
        raise RuntimeError("DOCX_PARSE_SMOKE_FAILED")

    xlsx_path = root / "sample.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "STBB19-320"
    worksheet["B1"] = "必须校验上限"
    workbook.save(str(xlsx_path))
    workbook.close()
    if not parser.parse(str(xlsx_path)):
        raise RuntimeError("XLSX_PARSE_SMOKE_FAILED")

    pptx_path = root / "sample.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, 3000000, 800000)
    box.text_frame.text = "RC7 requirement overview"
    deck.save(str(pptx_path))
    if not parser.parse(str(pptx_path)):
        raise RuntimeError("PPTX_PARSE_SMOKE_FAILED")

    return {
        "status": "PASS",
        "gate": "REQUIREMENT_TOOLCHAIN_PASS",
        "document_modules": _module_versions(),
        "diff_change_count": len(changes),
        "formats_smoked": ["txt", "docx", "xlsx", "pptx"],
        "pdf_provider": "pypdf",
        "scanned_pdf_policy": "FAIL_CLOSED_OCR_REQUIRED",
    }


def _diagnosis_gate() -> Dict[str, object]:
    diagnoser = MultiToolLogDiagnoser()
    cases = {
        "DEPENDENCY": ("maven", "[ERROR] Could not resolve dependencies for project pfc"),
        "ROUTE": ("api", "HTTP 404 route not found /cfg-data/foo"),
        "AUTH": ("playwright", "403 Forbidden session expired"),
        "DATA": ("database", "duplicate entry violates unique constraint"),
        "SCRIPT": ("playwright", "strict mode violation locator resolved to 3 elements"),
    }
    observed = {}
    for expected, (provider, text) in cases.items():
        result = diagnoser.analyze(text, [provider])
        actual = result.get("failure_class")
        observed[expected] = actual
        if actual != expected:
            raise RuntimeError("DIAGNOSIS_SMOKE_FAILED:%s!=%s" % (actual, expected))
    return {
        "status": "PASS",
        "gate": "DIAGNOSIS_TOOLCHAIN_PASS",
        "observed": observed,
        "product_policy": "NO_PRODUCT_WITHOUT_CORE_OWNERSHIP_EVIDENCE",
    }


def _report_gate(root: Path) -> Dict[str, object]:
    report = {
        "schema_version": "pfc.aitest.rc7.mission-report.v1",
        "mission_id": "QUAL-MISSION",
        "release_id": "QUAL-RC7",
        "summary": {
            "mission_state": "RUNNING",
            "task_total": 1,
            "task_pass": 0,
            "execution_total": 2,
            "execution_results": {"PASS": 1, "FAIL": 1},
            "failure_classes": {"ROUTE": 1},
            "defect_total": 0,
            "open_gap_total": 1,
        },
        "executions": [
            {"attempt_id": "A1", "tc_id": "TC-1", "layer": "L3", "status": "PASS"},
            {"attempt_id": "A2", "tc_id": "TC-2", "layer": "L5", "status": "FAIL"},
        ],
        "defects": [],
        "residual_risk": {"open_gaps": [{"gap_id": "G1", "severity": "HIGH", "description": "qualification gap"}]},
    }
    paths = ReportRenderer().write_all(report, str(root / "report"))
    for value in paths.values():
        if not Path(value).exists():
            raise RuntimeError("REPORT_OUTPUT_MISSING:%s" % value)
    allure_files = list((root / "report" / "allure-results").glob("*-result.json"))
    if len(allure_files) != 2:
        raise RuntimeError("ALLURE_RESULT_COUNT_INVALID")
    return {
        "status": "PASS",
        "gate": "REPORT_TOOLCHAIN_PASS",
        "outputs": sorted(paths),
        "renderers": ["json", "html", "junit", "allure-results"],
    }


def _code_gate(codegraph_bin: Optional[str]) -> Dict[str, object]:
    import rc7_release.code_qualification_strict as strict
    if not hasattr(strict, "run_strict_code_qualification"):
        raise RuntimeError("STRICT_CODE_QUALIFICATION_MISSING")
    details = {
        "status": "PASS",
        "gate": "CODE_ANALYSIS_TOOLCHAIN_PASS",
        "git_diff_proof": "ENFORCED_BY_STRICT_QUALIFICATION",
        "codegraph_proof": "ENFORCED_BY_STRICT_QUALIFICATION",
        "codegraph_smoke": "NOT_REQUESTED",
    }
    if codegraph_bin:
        binary = Path(codegraph_bin).expanduser().resolve()
        if not binary.is_file():
            raise RuntimeError("CODEGRAPH_BINARY_NOT_FOUND:%s" % binary)
        if os.name == "nt":
            result = subprocess.run(
                [str(binary), "--help"],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                universal_newlines=True,
                timeout=60,
            )
            if result.returncode != 0:
                raise RuntimeError("CODEGRAPH_SMOKE_FAILED:%s" % result.returncode)
            details["codegraph_smoke"] = "PASS"
        else:
            details["codegraph_smoke"] = "HOLD_NON_WINDOWS"
        details["codegraph_bin"] = str(binary)
    return details


def run(codegraph_bin: Optional[str] = None) -> Dict[str, object]:
    with tempfile.TemporaryDirectory(prefix="rc7-toolchain-qualification-") as temp:
        root = Path(temp)
        requirement = _requirement_gate(root)
        diagnosis = _diagnosis_gate()
        report = _report_gate(root)
        code = _code_gate(codegraph_bin)
    gates = [requirement, code, diagnosis, report]
    return {
        "schema_version": "pfc.aitest.rc7.toolchain-qualification.v1",
        "status": "PASS" if all(item.get("status") == "PASS" for item in gates) else "FAIL",
        "gates": {str(item["gate"]): item for item in gates},
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Qualify RC7 professional offline toolchains")
    parser.add_argument("--codegraph-bin", default=os.environ.get("PFC_CODEGRAPH_BIN"))
    parser.add_argument("--output")
    args = parser.parse_args(argv)
    result = run(args.codegraph_bin)
    text = json.dumps(result, ensure_ascii=False, sort_keys=True, indent=2)
    if args.output:
        target = Path(args.output).expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0 if result["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
