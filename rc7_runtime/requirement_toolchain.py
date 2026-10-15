from __future__ import annotations

import csv
import difflib
import json
import re
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, List, Optional, Sequence, Tuple


class RequirementToolchainError(RuntimeError):
    pass


@dataclass(frozen=True)
class SourceFragment:
    source_ref: str
    kind: str
    locator: str
    text: str

    def to_dict(self):
        return asdict(self)


@dataclass(frozen=True)
class RequirementDiff:
    change_type: str
    before: Optional[str]
    after: Optional[str]
    before_locator: Optional[str]
    after_locator: Optional[str]

    def to_dict(self):
        return asdict(self)


def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _fragments_from_lines(source_ref: str, kind: str, lines: Iterable[str]) -> List[SourceFragment]:
    out = []
    for index, line in enumerate(lines, 1):
        text = _normalize_text(line)
        if text:
            out.append(SourceFragment(source_ref, kind, "line:%d" % index, text))
    return out


class RequirementDocumentParser:
    """Offline-first document parser for requirement inputs.

    DOCX/XLSX/PPTX prefer installed document libraries when available and fall back
    to OOXML ZIP extraction. PDF uses pypdf and fails closed when no extractable
    text is present so scanned PDFs cannot be mistaken for parsed requirements.
    """

    TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".log"}

    def parse(self, path: str) -> Tuple[SourceFragment, ...]:
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise RequirementToolchainError("REQUIREMENT_SOURCE_NOT_FOUND:%s" % source)
        suffix = source.suffix.lower()
        if suffix in self.TEXT_SUFFIXES:
            return tuple(_fragments_from_lines(str(source), suffix.lstrip("."), source.read_text(encoding="utf-8-sig").splitlines()))
        if suffix == ".json":
            return tuple(self._parse_json(source))
        if suffix == ".csv":
            return tuple(self._parse_csv(source))
        if suffix == ".docx":
            return tuple(self._parse_docx(source))
        if suffix == ".xlsx":
            return tuple(self._parse_xlsx(source))
        if suffix == ".pptx":
            return tuple(self._parse_pptx(source))
        if suffix == ".pdf":
            return tuple(self._parse_pdf(source))
        raise RequirementToolchainError("UNSUPPORTED_REQUIREMENT_SOURCE:%s" % suffix)

    def _parse_json(self, source: Path) -> List[SourceFragment]:
        value = json.loads(source.read_text(encoding="utf-8-sig"))
        out: List[SourceFragment] = []

        def walk(node, locator: str) -> None:
            if isinstance(node, dict):
                for key, child in node.items():
                    walk(child, locator + "/" + str(key))
            elif isinstance(node, list):
                for index, child in enumerate(node):
                    walk(child, locator + "/%d" % index)
            else:
                text = _normalize_text(node)
                if text:
                    out.append(SourceFragment(str(source), "json", locator or "/", text))

        walk(value, "")
        return out

    def _parse_csv(self, source: Path) -> List[SourceFragment]:
        out = []
        with source.open("r", encoding="utf-8-sig", newline="") as fh:
            for r_index, row in enumerate(csv.reader(fh), 1):
                for c_index, value in enumerate(row, 1):
                    text = _normalize_text(value)
                    if text:
                        out.append(SourceFragment(str(source), "csv", "R%dC%d" % (r_index, c_index), text))
        return out

    def _parse_docx(self, source: Path) -> List[SourceFragment]:
        try:
            from docx import Document  # type: ignore
            doc = Document(str(source))
            out = []
            for index, paragraph in enumerate(doc.paragraphs, 1):
                text = _normalize_text(paragraph.text)
                if text:
                    out.append(SourceFragment(str(source), "docx", "paragraph:%d" % index, text))
            for t_index, table in enumerate(doc.tables, 1):
                for r_index, row in enumerate(table.rows, 1):
                    for c_index, cell in enumerate(row.cells, 1):
                        text = _normalize_text(cell.text)
                        if text:
                            out.append(SourceFragment(str(source), "docx", "table:%d/R%dC%d" % (t_index, r_index, c_index), text))
            return out
        except ImportError:
            return self._ooxml_text(source, "word/document.xml", "docx")

    def _parse_xlsx(self, source: Path) -> List[SourceFragment]:
        try:
            from openpyxl import load_workbook  # type: ignore
            workbook = load_workbook(str(source), read_only=True, data_only=True)
            out = []
            try:
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows():
                        for cell in row:
                            text = _normalize_text(cell.value)
                            if text:
                                out.append(SourceFragment(str(source), "xlsx", "%s!%s" % (sheet.title, cell.coordinate), text))
            finally:
                workbook.close()
            return out
        except ImportError:
            return self._ooxml_glob_text(source, "xl/worksheets/", "xlsx")

    def _parse_pptx(self, source: Path) -> List[SourceFragment]:
        try:
            from pptx import Presentation  # type: ignore
            presentation = Presentation(str(source))
            out = []
            for s_index, slide in enumerate(presentation.slides, 1):
                for sh_index, shape in enumerate(slide.shapes, 1):
                    text = _normalize_text(getattr(shape, "text", ""))
                    if text:
                        out.append(SourceFragment(str(source), "pptx", "slide:%d/shape:%d" % (s_index, sh_index), text))
            return out
        except ImportError:
            return self._ooxml_glob_text(source, "ppt/slides/", "pptx")

    def _parse_pdf(self, source: Path) -> List[SourceFragment]:
        try:
            from pypdf import PdfReader  # type: ignore
        except ImportError as exc:
            raise RequirementToolchainError("PDF_PROVIDER_NOT_AVAILABLE:pypdf") from exc
        reader = PdfReader(str(source))
        out = []
        for index, page in enumerate(reader.pages, 1):
            text = _normalize_text(page.extract_text() or "")
            if text:
                out.append(SourceFragment(str(source), "pdf", "page:%d" % index, text))
        if not out:
            raise RequirementToolchainError("OCR_REQUIRED:NO_EXTRACTABLE_PDF_TEXT")
        return out

    @staticmethod
    def _ooxml_text(source: Path, member: str, kind: str) -> List[SourceFragment]:
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(str(source), "r") as archive:
            raw = archive.read(member)
        root = ET.fromstring(raw)
        texts = [_normalize_text(node.text) for node in root.iter() if node.tag.endswith("}t") and node.text]
        return [SourceFragment(str(source), kind, "%s/text:%d" % (member, index), text)
                for index, text in enumerate(texts, 1) if text]

    @staticmethod
    def _ooxml_glob_text(source: Path, prefix: str, kind: str) -> List[SourceFragment]:
        import xml.etree.ElementTree as ET
        out = []
        with zipfile.ZipFile(str(source), "r") as archive:
            members = sorted(name for name in archive.namelist() if name.startswith(prefix) and name.endswith(".xml"))
            for member in members:
                root = ET.fromstring(archive.read(member))
                for index, node in enumerate((x for x in root.iter() if (x.tag.endswith("}t") or x.tag.endswith("}v")) and x.text), 1):
                    text = _normalize_text(node.text or "")
                    if text:
                        out.append(SourceFragment(str(source), kind, "%s/text:%d" % (member, index), text))
        return out


class RequirementDiffEngine:
    def diff(self, before: Sequence[SourceFragment], after: Sequence[SourceFragment]) -> Tuple[RequirementDiff, ...]:
        old = [fragment.text for fragment in before]
        new = [fragment.text for fragment in after]
        matcher = difflib.SequenceMatcher(a=old, b=new, autojunk=False)
        changes: List[RequirementDiff] = []
        for tag, i1, i2, j1, j2 in matcher.get_opcodes():
            if tag == "equal":
                continue
            if tag == "replace":
                span = max(i2 - i1, j2 - j1)
                for offset in range(span):
                    b = before[i1 + offset] if i1 + offset < i2 else None
                    a = after[j1 + offset] if j1 + offset < j2 else None
                    changes.append(RequirementDiff(
                        "MODIFIED",
                        b.text if b else None,
                        a.text if a else None,
                        b.locator if b else None,
                        a.locator if a else None,
                    ))
            elif tag == "delete":
                for fragment in before[i1:i2]:
                    changes.append(RequirementDiff("REMOVED", fragment.text, None, fragment.locator, None))
            elif tag == "insert":
                for fragment in after[j1:j2]:
                    changes.append(RequirementDiff("ADDED", None, fragment.text, None, fragment.locator))
        return tuple(changes)


def source_map(fragments: Sequence[SourceFragment]) -> dict:
    return {
        "fragment_count": len(fragments),
        "sources": sorted({fragment.source_ref for fragment in fragments}),
        "fragments": [fragment.to_dict() for fragment in fragments],
    }
