from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import time
from dataclasses import asdict
from pathlib import Path

from docx import Document

from rc6_runtime.automation import PytestRunnerAdapter, ScriptArtifact
from rc6_runtime.diagnosis import DiagnosisResult, RecoveryDecision
from rc6_runtime.execution import ExecutionRecord
from rc6_runtime.truth import TruthRuntime
from rc7_runtime import (
    Artifact,
    ClosedLoopControlPlane,
    Evaluation,
    EvaluationResult,
    Evidence,
    Mission,
    MissionTaskGraphCompiler,
    TaskStatus,
)
from rc7_runtime.code_loop import CodeAnalysisClosedLoop
from rc7_runtime.control_plane import WorkerOutput
from rc7_runtime.defect_loop import (
    DefectClosedLoop,
    DefectQualificationInput,
    DefectStatus,
    RootCauseHypothesis,
)
from rc7_runtime.execution_loop import ExecutionClosedLoop
from rc7_runtime.mission_evaluation_runtime import MissionEvaluationRuntime
from rc7_runtime.reporting import MissionReportBuilder, ReportRenderer
from rc7_runtime.requirement_loop import RequirementClosedLoop
from rc7_runtime.script_readiness_loop import ScriptReadinessClosedLoop
from rc7_runtime.test_design_loop import TestDesignClosedLoop, TestObligation
from rc7_runtime.truth_overlay import append_entity

MISSION = "MISSION-RC7-REAL-DEMO-001"
RELEASE = "RC7-REAL-DEMO"
REQ = "REQ-RATE-001"
TC = "TC-RATE-BOUNDARY-001"
TARGET = "demo-rate-service"


def run(command, *, cwd=None, check=True, env=None):
    completed = subprocess.run(
        [str(x) for x in command], cwd=str(cwd) if cwd else None,
        env=env, text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        check=False,
    )
    if check and completed.returncode != 0:
        raise RuntimeError("COMMAND_FAILED:%s\nSTDOUT:\n%s\nSTDERR:\n%s" % (
            " ".join(str(x) for x in command), completed.stdout, completed.stderr
        ))
    return completed


def sha256_path(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, payload) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2, default=str) + "\n"
    path.write_text(encoded, encoding="utf-8")
    return "file://" + str(path.resolve())


def make_requirement(source_md: Path, out: Path):
    md = source_md.read_text(encoding="utf-8")
    copied = out / "inputs" / source_md.name
    copied.parent.mkdir(parents=True, exist_ok=True)
    copied.write_text(md, encoding="utf-8")
    doc = Document()
    doc.add_heading("REQ-RATE-001 利率调整边界规则", level=0)
    doc.add_paragraph("本需求书由演示流程从受版本控制的 Markdown 原始需求生成，用于验证真实文档输入链路。")
    for line in md.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("# REQ-"):
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        else:
            doc.add_paragraph(stripped)
    docx = out / "inputs" / "REQ-RATE-001.docx"
    doc.save(str(docx))
    return copied, docx


def init_target_repo(out: Path):
    repo = out / "target-repo"
    if repo.exists():
        shutil.rmtree(repo)
    repo.mkdir(parents=True)
    run(["git", "init", "-b", "main"], cwd=repo)
    run(["git", "config", "user.email", "rc7-demo@example.invalid"], cwd=repo)
    run(["git", "config", "user.name", "RC7 Demo"], cwd=repo)
    good = '''from decimal import Decimal, ROUND_HALF_UP\n\ndef adjust_rate(current_rate: str, delta: str):\n    current = Decimal(current_rate)\n    change = Decimal(delta)\n    if change < Decimal("-0.50") or change > Decimal("0.50"):\n        return {"accepted": False, "code": "DELTA_OUT_OF_RANGE"}\n    result = current + change\n    if result < 0:\n        return {"accepted": False, "code": "NEGATIVE_RATE_FORBIDDEN"}\n    return {"accepted": True, "code": "ACCEPTED", "rate": str(result.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))}\n'''
    buggy = good.replace('change > Decimal("0.50")', 'change >= Decimal("0.50")')
    service = repo / "rate_service.py"
    service.write_text(good, encoding="utf-8")
    (repo / "README.md").write_text("# Demo Rate Service\n\nTarget repository for RC7 real-mission qualification.\n", encoding="utf-8")
    run(["git", "add", "."], cwd=repo)
    run(["git", "commit", "-m", "feat: implement inclusive rate adjustment boundaries"], cwd=repo)
    base = run(["git", "rev-parse", "HEAD"], cwd=repo).stdout.strip()
    run(["git", "checkout", "-b", "feature/rate-boundary-change"], cwd=repo)
    service.write_text(buggy, encoding="utf-8")
    run(["git", "add", "rate_service.py"], cwd=repo)
    run(["git", "commit", "-m", "refactor: adjust rate boundary predicate"], cwd=repo)
    buggy_sha = run(["git", "rev-parse", "HEAD"], cwd=repo).stdout.strip()
    return repo, service, good, base, buggy_sha


def codegraph_probe(binary: Path, repo: Path, out: Path):
    args = [
        str(binary), "--workspace", str(repo), "--graph-only",
        "--run-tool", "codegraph_pr_context",
        "--tool-args", json.dumps({"baseBranch": "main", "format": "markdown"}),
    ]
    completed = run(args, cwd=repo, check=False)
    path = out / "evidence" / "codegraph-pr-context.txt"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "COMMAND: %s\nEXIT_CODE: %s\n\nSTDOUT:\n%s\n\nSTDERR:\n%s\n" % (
            " ".join(args), completed.returncode, completed.stdout, completed.stderr
        ), encoding="utf-8"
    )
    if completed.returncode != 0 or not completed.stdout.strip():
        raise RuntimeError("CODEGRAPH_REAL_PROBE_FAILED")
    return path


def persist_candidate(out: Path, artifact_id: str, payload):
    path = out / "node-artifacts" / (artifact_id + ".json")
    ref = write_json(path, payload)
    return ref, "sha256:" + sha256_path(path)


def produce(control, graph, compiler, node: str, worker: str, out: Path, payload):
    task = graph.task(node)
    leased = control.lease_next(worker_ref=worker, capabilities=compiler.REQUIRED_CAPABILITIES[node])
    if leased is None or leased.task_id != task.task_id:
        raise RuntimeError("UNEXPECTED_TASK_LEASE:%s" % node)
    control.start(task.task_id, worker_ref=worker)
    evidence_id = "EVD-WORKER-%s" % node
    artifact_id = graph.artifact_ids[node]
    content_ref, content_hash = persist_candidate(out, artifact_id, payload)
    evidence_kind = "NODE_OUTPUT_PROOF"
    evidence_source = "worker://%s" % worker
    subjects = [task.task_id, artifact_id]
    if node == "EXECUTION":
        record = payload["record"]
        evidence_kind = "RUNTIME_EXECUTION_RECORD"
        evidence_source = "runtime://pytest-runner"
        subjects.extend([record["attempt_id"], record["tc_id"]])
    candidate = control.submit_worker_output(WorkerOutput(
        task_id=task.task_id,
        worker_ref=worker,
        evidence=(Evidence(
            evidence_id=evidence_id, mission_id=MISSION, kind=evidence_kind,
            source_ref=evidence_source, content_ref=content_ref,
            subject_refs=subjects, content_hash=content_hash,
        ),),
        artifacts=(Artifact(
            artifact_id=artifact_id, mission_id=MISSION,
            artifact_type=compiler.OUTPUT_TYPES[node], producer_task_id=task.task_id,
            content_ref=content_ref, parent_refs=list(task.input_artifact_refs),
            evidence_refs=[evidence_id], content_hash=content_hash,
        ),),
    ))
    if candidate.state != TaskStatus.OUTPUT_PRODUCED:
        raise RuntimeError("OUTPUT_NOT_PRODUCED:%s" % node)
    return task, artifact_id


def produce_dynamic(control, task, worker: str, out: Path, artifact_id: str, artifact_type: str, payload, evidence_id: str):
    leased = control.lease_next(worker_ref=worker, capabilities=task.required_capabilities)
    if leased is None or leased.task_id != task.task_id:
        raise RuntimeError("UNEXPECTED_DYNAMIC_TASK_LEASE:%s" % task.task_id)
    control.start(task.task_id, worker_ref=worker)
    content_ref, content_hash = persist_candidate(out, artifact_id, payload)
    candidate = control.submit_worker_output(WorkerOutput(
        task_id=task.task_id,
        worker_ref=worker,
        evidence=(Evidence(
            evidence_id=evidence_id, mission_id=MISSION, kind="RUNTIME_EXECUTION_RECORD",
            source_ref="runtime://pytest-runner", content_ref=content_ref,
            subject_refs=[task.task_id, artifact_id], content_hash=content_hash,
        ),),
        artifacts=(Artifact(
            artifact_id=artifact_id, mission_id=MISSION, artifact_type=artifact_type,
            producer_task_id=task.task_id, content_ref=content_ref,
            parent_refs=list(task.input_artifact_refs), evidence_refs=[evidence_id],
            content_hash=content_hash,
        ),),
    ))
    if candidate.state != TaskStatus.OUTPUT_PRODUCED:
        raise RuntimeError("DYNAMIC_OUTPUT_NOT_PRODUCED")
    return content_ref


def append_business_truth(truth: TruthRuntime, event_type: str, payload):
    truth.append_atomic(event_type, payload)


def execution_payload(record: ExecutionRecord, status: str):
    return {
        "attempt_id": record.attempt_id,
        "release_id": RELEASE,
        "mission_id": MISSION,
        "requirement_id": REQ,
        "tc_id": record.tc_id,
        "case_version": record.case_version,
        "target_id": record.target_id,
        "layer": record.layer,
        "status": status,
        "test_result": status,
        "actual_result": record.actual_result,
        "tool_call_refs": list(record.tool_call_refs),
        "evidence_refs": list(record.evidence_refs),
        "reason": record.reason,
    }


def make_record(attempt_id: str, runner_result, *, passed: bool):
    return ExecutionRecord(
        attempt_id=attempt_id, tc_id=TC, case_version=1, target_id=TARGET,
        layer="L3", worker_id="worker-execution", lease_id="lease-%s" % attempt_id,
        status="PASS" if passed else "FAIL", integrity_passed=True,
        actual_result=dict(runner_result.actual_result),
        tool_call_refs=tuple(runner_result.tool_call_refs),
        evidence_refs=tuple(runner_result.evidence_refs),
        assertion_ref="oracle://REQ-RATE-001/upper-bound-inclusive",
        assertion_passed=passed,
        reason="ORACLE_PASS" if passed else "ORACLE_ASSERTION_FAILED",
    )


def build_script(repo: Path, spec):
    source = repo / "test_req_rate_001.py"
    source.write_text(r'''import json
import os
from rate_service import adjust_rate

def _emit(row):
    path = os.environ.get("RC6_EXECUTION_MANIFEST")
    if not path:
        return
    with open(path, "a", encoding="utf-8") as fh:
        fh.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")

def _call(current, delta):
    result = adjust_rate(current, delta)
    _emit({"type":"tool_call","tool_id":"api-runtime","call_ref":"python://rate_service.adjust_rate","evidence_ref":"runtime://rate-service-call"})
    _emit({"type":"actual_result","value":result})
    return result

def test_upper_boundary_accepted():
    result = _call("3.25", "0.50")
    assert result["accepted"] is True
    assert result["code"] == "ACCEPTED"
    assert result["rate"] == "3.75"

def test_above_boundary_rejected():
    result = _call("3.25", "0.51")
    assert result == {"accepted": False, "code": "DELTA_OUT_OF_RANGE"}

def test_negative_result_forbidden():
    result = _call("0.20", "-0.50")
    assert result == {"accepted": False, "code": "NEGATIVE_RATE_FORBIDDEN"}
''', encoding="utf-8")
    return source, ScriptArtifact(
        script_id="SCRIPT-RATE-REAL-DEMO", spec_id=spec.spec_id, tc_id=spec.tc_id,
        case_version=spec.case_version, layer=spec.layer, runner_type=spec.runner_type,
        source_ref="file://" + str(source.resolve()), test_symbol="test_upper_boundary_accepted",
        case_fingerprint=spec.case_fingerprint, oracle_fingerprint=spec.oracle_fingerprint,
        reuse_refs=(), builder_evidence_refs=("file://" + str(source.resolve()),), revision=1,
    )


def persist_defect_version(truth, defect, version):
    DefectClosedLoop.persist(truth, release_id=RELEASE, defect=defect, artifact_version=version)


def main(argv=None):
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True)
    parser.add_argument("--codegraph", required=True)
    args = parser.parse_args(argv)
    out = Path(args.out).resolve()
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    source_md = Path(__file__).resolve().parent / "REQ-RATE-001.md"
    req_md, req_docx = make_requirement(source_md, out)
    repo, service, fixed_source, base_sha, buggy_sha = init_target_repo(out)
    diff_path = out / "evidence" / "git-diff-base-to-buggy.patch"
    diff_path.parent.mkdir(parents=True, exist_ok=True)
    diff = run(["git", "diff", "%s..%s" % (base_sha, buggy_sha)], cwd=repo).stdout
    diff_path.write_text(diff, encoding="utf-8")
    if '>= Decimal("0.50")' not in diff:
        raise RuntimeError("EXPECTED_REAL_CODE_CHANGE_NOT_FOUND")
    codegraph_path = codegraph_probe(Path(args.codegraph).resolve(), repo, out)

    r1 = out / "state" / "rc7-r1.jsonl"
    truth = TruthRuntime(r1)
    control = ClosedLoopControlPlane.bootstrap(
        truth, Mission(mission_id=MISSION, release_id=RELEASE, goal="真实需求书+真实Git差异+真实pytest执行+缺陷闭环+报告+后管")
    )
    compiler = MissionTaskGraphCompiler()
    graph = compiler.install(control, MISSION, external_inputs={
        "REQUIREMENT_ANALYSIS": ("file://" + str(req_docx),),
        "CODE_ANALYSIS": ("git://%s..%s" % (base_sha, buggy_sha),),
    })

    req_analysis = {
        "analysis_id": "REQ-AN-REAL-DEMO", "release_id": RELEASE, "requirement_id": REQ,
        "analyzer_worker_id": "worker-requirement",
        "source_assertions": [{"id":"SRC-BOUNDARY","text":"delta范围[-0.50,+0.50]包含边界"}],
        "atomic_statements": [{"id":"A-UPPER","statement":"delta=+0.50必须受理并得到3.75","testable":True,"source_ref":"SRC-BOUNDARY"}],
        "trace_links": [{"source_id":"SRC-BOUNDARY","atomic_id":"A-UPPER","evidence_ref":"file://" + str(req_docx)}],
        "evidence_refs": ["file://" + str(req_docx)],
        "business_rules": [{"id":"BR-RANGE"}], "roles": [{"id":"ROLE-LOAN-OPS"}],
        "preconditions": [{"id":"PRE-RATE-VALID"}], "normal_paths": [{"id":"NP-ADJUST"}],
        "negative_paths": [{"id":"NEG-OUTSIDE"},{"id":"NEG-BELOW-ZERO"}],
        "boundaries": [{"id":"BOUND-UPPER","rule":"+0.50 inclusive"},{"id":"BOUND-ABOVE","rule":"+0.51 rejected"}],
        "states": [{"id":"STATE-ACCEPTED"},{"id":"STATE-REJECTED"}],
        "data_constraints": [{"id":"DATA-DECIMAL-2"}], "dependencies": [{"id":"DEP-RATE-SERVICE"}],
        "side_effects": [{"id":"SIDE-NONE-DEMO"}],
    }
    req_review = {
        "id":"REQ-REVIEW-REAL-DEMO","review_id":"REQ-REVIEW-REAL-DEMO","release_id":RELEASE,
        "requirement_id":REQ,"subject_analysis_id":req_analysis["analysis_id"],
        "reviewer_worker_id":"evaluator-requirement","analyzer_worker_id":"worker-requirement",
        "decision":"PASS","evidence_refs":["file://" + str(req_md)],"gap_ids":[],
    }
    req_task, req_art = produce(control, graph, compiler, "REQUIREMENT_ANALYSIS", "worker-requirement", out, {"analysis":req_analysis,"review":req_review})
    req_result = RequirementClosedLoop().evaluate(mission_id=MISSION, artifact_ref=req_art, analysis=req_analysis, review=req_review)
    if req_result.evaluation.result != EvaluationResult.PASS:
        raise RuntimeError("REAL_REQUIREMENT_GATE_FAILED:%s" % req_result.evaluation.dimensions)
    control.accept_evaluation(req_task.task_id, evaluation=req_result.evaluation)
    append_business_truth(truth, "REQUIREMENT_ANALYSIS_UPSERTED", {
        "id":"REQ-AN-REAL-DEMO","release_id":RELEASE,"mission_id":MISSION,"requirement_id":REQ,
        "status":"PASS","br":["BR-RANGE","BR-NONNEGATIVE"],"sr":["SR-INCLUSIVE-UPPER"],"tr":["TR-BOUNDARY-ORACLE"],
        "evidence_refs":["file://" + str(req_docx)],
    })

    code_analysis = {
        "analysis_id":"CODE-AN-REAL-DEMO","release_id":RELEASE,"analyzer_worker_id":"worker-code",
        "requirement_ids":[REQ],"source_revisions":{TARGET:"%s..%s"%(base_sha,buggy_sha)},
        "changed_files":["rate_service.py"],
        "changed_symbols":[{"id":"SYM-ADJUST","symbol":"rate_service.adjust_rate"}],
        "symbol_explanations":[{"symbol_id":"SYM-ADJUST","evidence_refs":["file://" + str(diff_path),"file://" + str(codegraph_path)]}],
        "reverse_requirement_mappings":[{"requirement_id":REQ,"symbol_id":"SYM-ADJUST","evidence_ref":"file://" + str(diff_path)}],
        "tool_calls":[
            {"tool_id":"git","call_ref":"git://%s..%s"%(base_sha,buggy_sha),"evidence_ref":"file://" + str(diff_path),"target_id":TARGET},
            {"tool_id":"codegraph","call_ref":"codegraph://codegraph_pr_context","evidence_ref":"file://" + str(codegraph_path),"target_id":TARGET},
        ],
        "evidence_refs":["file://" + str(diff_path),"file://" + str(codegraph_path)],"codegraph_required":True,
    }
    code_review = {
        "review_id":"CODE-REVIEW-REAL-DEMO","subject_analysis_id":code_analysis["analysis_id"],
        "reviewer_worker_id":"evaluator-code","analyzer_worker_id":"worker-code","decision":"PASS",
        "evidence_refs":["file://" + str(codegraph_path)],"gap_ids":[],
    }
    code_task, code_art = produce(control, graph, compiler, "CODE_ANALYSIS", "worker-code", out, {"analysis":code_analysis,"review":code_review})
    code_result = CodeAnalysisClosedLoop().evaluate(mission_id=MISSION, artifact_ref=code_art, target_id=TARGET, analysis=code_analysis, review=code_review)
    if code_result.evaluation.result != EvaluationResult.PASS:
        raise RuntimeError("REAL_CODE_GATE_FAILED:%s" % code_result.evaluation.dimensions)
    control.accept_evaluation(code_task.task_id, evaluation=code_result.evaluation)
    append_business_truth(truth, "CODE_ANALYSIS_UPSERTED", {
        "id":"CODE-AN-REAL-DEMO","release_id":RELEASE,"mission_id":MISSION,"requirement_id":REQ,"status":"PASS",
        "revision_refs":[base_sha,buggy_sha],"impact_refs":["rate_service.py#adjust_rate"],
        "evidence_refs":["file://" + str(diff_path),"file://" + str(codegraph_path)],
    })

    obligations = [
        TestObligation("OBL-BR-RANGE","BUSINESS_RULE","L3","HIGH",("file://" + str(req_docx),)),
        TestObligation("OBL-UPPER-INCLUSIVE","BOUNDARY","L3","CRITICAL",("file://" + str(req_docx),)),
        TestObligation("OBL-CODE-PREDICATE","CODE_IMPACT","L3","CRITICAL",("file://" + str(diff_path),)),
    ]
    design_cases = [{
        "tc_id":TC,"layer":"L3","risk":"CRITICAL","status":"CANDIDATE",
        "coverage_links":[{"subject_id":o.obligation_id,"evidence_ref":"coverage://%s"%o.obligation_id} for o in obligations],
        "action_semantics":{"state_changing":False,"negative":False,"authorization_sensitive":False},
        "oracle":{"checks":[
            {"type":"API_RESPONSE","evidence_ref":"oracle://accepted"},
            {"type":"BUSINESS_RULE","evidence_ref":"oracle://inclusive-upper"},
            {"type":"EVIDENCE_CORRELATION","evidence_ref":"oracle://runtime-manifest"},
        ]},
    }]
    design_payload={"artifact":{"artifact_id":graph.artifact_ids["TEST_DESIGN"],"artifact_type":"TEST_DESIGN_MODEL"},"semantic_payload":{"obligations":[asdict(x) for x in obligations],"cases":design_cases}}
    design_task, design_art = produce(control, graph, compiler, "TEST_DESIGN", "worker-design", out, design_payload)
    design_result = TestDesignClosedLoop().evaluate(mission_id=MISSION, artifact_ref=design_art, obligations=obligations, cases=design_cases)
    if design_result.evaluation.result != EvaluationResult.PASS or design_result.semantic_coverage != 1.0:
        raise RuntimeError("REAL_DESIGN_GATE_FAILED")
    control.accept_evaluation(design_task.task_id, evaluation=design_result.evaluation)

    frozen_case={
        "tc_id":TC,"case_version":1,"release_id":RELEASE,"requirement_id":REQ,"sst_id":"SST-DEMO-RATE-001",
        "condition_id":"COND-UPPER-INCLUSIVE","objective":"验证+0.50上边界包含规则","layer":"L3","status":"FROZEN",
        "preconditions":["rate service importable"],"test_data":{"current_rate":"3.25","delta":"0.50"},
        "steps":["调用adjust_rate(3.25,0.50)"],"expected":"accepted=true, code=ACCEPTED, rate=3.75",
        "oracle":{"checks":[{"type":"BUSINESS_RULE","expected":"inclusive upper boundary"}]},
        "evidence_requirements":["pytest process","runtime manifest","actual result"],
    }
    append_business_truth(truth,"TEST_CASE_UPSERTED",dict(frozen_case,mission_id=MISSION))
    script_loop=ScriptReadinessClosedLoop(validation_timeout_seconds=60)
    spec=script_loop.compile_execution_contract(
        case=frozen_case,target_requirements=(REQ,),test_data_refs=("inline://rate-boundary-data",),
        oracle_refs=("oracle://REQ-RATE-001/upper-bound-inclusive",),cleanup_policy={"required":False},
        runner_type="PYTEST",automation_mapping={"reuse":{}},
    )
    source, script=build_script(repo,spec)
    validation=script_loop.validate_script(spec=spec,script=script,cwd=str(repo),env={"PYTHONPATH":str(repo)})
    script_task, script_art=produce(control,graph,compiler,"SCRIPT_READINESS","worker-script",out,{"spec":asdict(spec),"script":asdict(script),"validation":asdict(validation)})
    script_result=script_loop.evaluate(mission_id=MISSION,artifact_ref=script_art,spec=spec,script=script,validation=validation)
    if script_result.evaluation.result != EvaluationResult.PASS:
        raise RuntimeError("SCRIPT_NOT_READY")
    control.accept_evaluation(script_task.task_id,evaluation=script_result.evaluation)

    adapter=PytestRunnerAdapter(timeout_seconds=60)
    run1=adapter.execute(spec=spec,script=script,validation=validation,cwd=str(repo),env={"PYTHONPATH":str(repo)})
    if run1.exit_code == 0 or run1.actual_result.get("code") != "DELTA_OUT_OF_RANGE":
        raise RuntimeError("KNOWN_PRODUCT_BUG_NOT_OBSERVED")
    rec1=make_record("ATT-RATE-001",run1,passed=False)
    append_business_truth(truth,"EXECUTION_ATTEMPT_UPSERTED",execution_payload(rec1,"FAIL"))
    diag=DiagnosisResult(
        diagnosis_id="DIAG-RATE-001",attempt_id=rec1.attempt_id,failure_class="PRODUCT",
        evidence_refs=rec1.evidence_refs,cat_required=False,cat_queried=False,
        rationale="需求规定+0.50必须受理；脚本已就绪且真实pytest执行返回DELTA_OUT_OF_RANGE，代码diff显示>=0.50边界谓词。",
        classifier_ref="demo-evidence-backed-classifier",
    )
    recovery=RecoveryDecision(
        diagnosis_id=diag.diagnosis_id,attempt_id=rec1.attempt_id,action="PRODUCT_DEFECT",
        retry_allowed=False,reprepare_allowed=False,terminal=True,next_attempt_id=None,
        reason="trusted business oracle mismatch",
    )
    exec_task,exec_art=produce(control,graph,compiler,"EXECUTION","worker-execution",out,{"record":asdict(rec1),"diagnosis":asdict(diag),"recovery":asdict(recovery)})
    exec_result=ExecutionClosedLoop().evaluate(mission_id=MISSION,artifact_ref=exec_art,record=rec1,diagnosis=diag,recovery=recovery)
    if exec_result.test_result!="FAIL" or exec_result.evaluation.result!=EvaluationResult.PASS or not exec_result.gaps:
        raise RuntimeError("PRODUCT_DEFECT_HANDOFF_NOT_CREATED")
    control.accept_evaluation(exec_task.task_id,evaluation=exec_result.evaluation,gaps=exec_result.gaps)
    product_gap=exec_result.gaps[0]

    repro_rows=[r for r in truth.query("rc7_task",mission_id=MISSION) if r.get("task_type")=="DEFECT_REPRODUCTION"]
    if len(repro_rows)!=1: raise RuntimeError("DEFECT_REPRO_TASK_MISSING")
    repro_task=control.scheduler.task(repro_rows[0]["task_id"])
    run2=adapter.execute(spec=spec,script=script,validation=validation,cwd=str(repo),env={"PYTHONPATH":str(repo)})
    if run2.exit_code==0 or run2.actual_result != run1.actual_result:
        raise RuntimeError("DEFECT_NOT_REPRODUCED")
    rec2=make_record("ATT-RATE-002-REPRO",run2,passed=False)
    append_business_truth(truth,"EXECUTION_ATTEMPT_UPSERTED",execution_payload(rec2,"FAIL"))
    repro_art="ART-DEFECT-REPRO-RATE-001"
    produce_dynamic(control,repro_task,"worker-defect-repro",out,repro_art,"DEFECT_REPRODUCTION_EVIDENCE",{"record":asdict(rec2),"same_actual":True},"EVD-DEFECT-REPRO")
    repro_eval=Evaluation(
        evaluation_id="EVAL-DEFECT-REPRO-RATE-001",mission_id=MISSION,node="EXECUTION",
        artifact_refs=[repro_art],result=EvaluationResult.PASS,evaluator_ref="evaluator-defect-reproduction",
        dimensions={"same_failure_reproduced":True,"attempt_id":rec2.attempt_id,"test_result":"FAIL"},
        gap_refs=[],evidence_refs=list(rec2.evidence_refs),
    )
    control.accept_evaluation(repro_task.task_id,evaluation=repro_eval)

    qual_rows=[r for r in truth.query("rc7_task",mission_id=MISSION) if r.get("task_type")=="QUALIFY_DEFECT"]
    if len(qual_rows)!=1: raise RuntimeError("DEFECT_QUALIFY_TASK_MISSING")
    qual_task=control.scheduler.task(qual_rows[0]["task_id"])
    leased=control.lease_next(worker_ref="worker-defect",capabilities=qual_task.required_capabilities)
    if leased is None or leased.task_id!=qual_task.task_id: raise RuntimeError("DEFECT_QUALIFY_LEASE_FAILED")
    control.start(qual_task.task_id,worker_ref="worker-defect")

    defect_loop=DefectClosedLoop()
    defect=defect_loop.create_candidate(
        mission_id=MISSION,product_gap=product_gap,requirement_refs=[REQ],test_case_id=TC,
        execution_attempt_id=rec1.attempt_id,expected={"accepted":True,"code":"ACCEPTED","rate":"3.75"},
        actual=rec1.actual_result,evidence_refs=list(rec1.evidence_refs),severity="HIGH",
    )
    defect_loop.transition(defect,DefectStatus.QUALIFYING)
    defect_loop.record_reproduction(defect,attempt_id=rec2.attempt_id,reproduced=True,evidence_refs=list(rec2.evidence_refs))
    qual_result=defect_loop.evaluate_qualification(defect,DefectQualificationInput(
        execution_trusted=True,oracle_valid=True,evidence_complete=True,script_issue_excluded=True,
        data_issue_excluded=True,environment_issue_excluded=True,reproducible=True,
    ))
    if not qual_result.qualified: raise RuntimeError("DEFECT_QUALIFICATION_FAILED")
    append_entity(truth,qual_result.evaluation,release_id=RELEASE,artifact_version=1)
    persist_defect_version(truth,defect,1)
    hypothesis=RootCauseHypothesis(
        hypothesis_id="RCA-RATE-UPPER-BOUND",statement="rate_service.adjust_rate将+0.50误写为排除边界(>=)，导致合法请求被拒绝。",
        suspected_symbols=("rate_service.adjust_rate",),evidence_refs=("file://"+str(diff_path),"file://"+str(codegraph_path)),confidence=0.99,
    )
    defect_loop.attach_root_cause_hypothesis(defect,hypothesis)
    defect_loop.transition(defect,DefectStatus.READY_FOR_REVIEW)
    persist_defect_version(truth,defect,2)
    defect_loop.transition(defect,DefectStatus.CONFIRMED,evaluation_id=qual_result.evaluation.evaluation_id)
    persist_defect_version(truth,defect,3)
    defect_loop.transition(defect,DefectStatus.FIX_PENDING)
    persist_defect_version(truth,defect,4)

    service.write_text(fixed_source,encoding="utf-8")
    run(["git","add","rate_service.py"],cwd=repo)
    run(["git","commit","-m","fix: restore inclusive upper rate boundary"],cwd=repo)
    fix_sha=run(["git","rev-parse","HEAD"],cwd=repo).stdout.strip()
    fix_diff=out/"evidence"/"git-diff-buggy-to-fix.patch"
    fix_diff.write_text(run(["git","diff","%s..%s"%(buggy_sha,fix_sha)],cwd=repo).stdout,encoding="utf-8")
    defect_loop.transition(defect,DefectStatus.FIX_DETECTED,fix_revision=fix_sha)
    persist_defect_version(truth,defect,5)
    defect_loop.transition(defect,DefectStatus.VERIFYING)
    persist_defect_version(truth,defect,6)

    run3=adapter.execute(spec=spec,script=script,validation=validation,cwd=str(repo),env={"PYTHONPATH":str(repo)})
    if run3.exit_code!=0 or run3.actual_result.get("rate")!="3.75": raise RuntimeError("FIX_VERIFICATION_FAILED")
    rec3=make_record("ATT-RATE-003-FIX-VERIFY",run3,passed=True)
    append_business_truth(truth,"EXECUTION_ATTEMPT_UPSERTED",execution_payload(rec3,"PASS"))
    verify_ref="runtime://"+rec3.attempt_id
    defect_loop.transition(defect,DefectStatus.FIXED,verification_refs=[verify_ref])
    persist_defect_version(truth,defect,7)
    defect_loop.transition(defect,DefectStatus.REGRESSION_TESTING)
    persist_defect_version(truth,defect,8)

    regression=run([sys.executable,"-m","pytest","-q",str(source)],cwd=repo,check=False,env=dict(os.environ,PYTHONPATH=str(repo)))
    regression_log=out/"evidence"/"regression-full-suite.txt"
    regression_log.write_text("EXIT_CODE=%s\nSTDOUT:\n%s\nSTDERR:\n%s"%(regression.returncode,regression.stdout,regression.stderr),encoding="utf-8")
    if regression.returncode!=0 or "3 passed" not in regression.stdout: raise RuntimeError("REGRESSION_FAILED")
    rec4=ExecutionRecord(
        attempt_id="ATT-RATE-004-REGRESSION",tc_id=TC,case_version=1,target_id=TARGET,layer="L3",
        worker_id="worker-regression",lease_id="lease-regression",status="PASS",integrity_passed=True,
        actual_result={"suite":"3 passed","fix_revision":fix_sha},tool_call_refs=("pytest://full-regression",),
        evidence_refs=("file://"+str(regression_log),),assertion_ref="oracle://regression-suite",assertion_passed=True,reason="REGRESSION_PASS",
    )
    append_business_truth(truth,"EXECUTION_ATTEMPT_UPSERTED",execution_payload(rec4,"PASS"))

    close_eval_id="EVAL-DEFECT-CLOSE-RATE-001"
    defect_loop.transition(defect,DefectStatus.CLOSED,evaluation_id=close_eval_id,regression_refs=["file://"+str(regression_log)])
    defect_json=out/"node-artifacts"/(defect.defect_id+".json")
    content_ref=write_json(defect_json,asdict(defect))
    content_hash="sha256:"+sha256_path(defect_json)
    control.submit_worker_output(WorkerOutput(
        task_id=qual_task.task_id,worker_ref="worker-defect",
        evidence=(Evidence(
            evidence_id="EVD-DEFECT-CLOSED",mission_id=MISSION,kind="DEFECT_LIFECYCLE_PROOF",
            source_ref="runtime://defect-closed-loop",content_ref=content_ref,
            subject_refs=[qual_task.task_id,defect.defect_id],content_hash=content_hash,
        ),),
        artifacts=(Artifact(
            artifact_id=defect.defect_id,mission_id=MISSION,artifact_type="DEFECT_LIFECYCLE_CONCLUSION",
            producer_task_id=qual_task.task_id,content_ref=content_ref,parent_refs=list(qual_task.input_artifact_refs),
            evidence_refs=["EVD-DEFECT-CLOSED"],content_hash=content_hash,
        ),),
    ))
    close_eval=Evaluation(
        evaluation_id=close_eval_id,mission_id=MISSION,node="DEFECT",artifact_refs=[defect.defect_id],
        result=EvaluationResult.PASS,evaluator_ref="evaluator-defect-terminal",
        dimensions={"defect_id":defect.defect_id,"defect_state":"CLOSED","fix_revision":fix_sha,"verification":True,"regression":True},
        gap_refs=[],evidence_refs=[verify_ref,"file://"+str(regression_log)],
    )
    control.accept_defect_resolution(qual_task.task_id,defect=defect,evaluation=close_eval)

    mission_result=MissionEvaluationRuntime(evidence_root=str(out/"mission-evidence")).run(control)
    if mission_result.evaluation_result!="PASS" or mission_result.mission_state!="CONVERGED":
        raise RuntimeError("MISSION_DID_NOT_CONVERGE:%s"%mission_result)

    builder=MissionReportBuilder(truth)
    report=builder.build(MISSION)
    report["summary"]["release_recommendation"]=mission_result.release_recommendation
    report["summary"]["real_product_defect_found_and_closed"]=True
    report["demo_truth"]={
        "requirement_docx":"file://"+str(req_docx),"target_repo":str(repo),"base_commit":base_sha,
        "buggy_commit":buggy_sha,"fix_commit":fix_sha,"codegraph_evidence":"file://"+str(codegraph_path),
        "initial_test_result":"FAIL","reproduction_test_result":"FAIL","fix_verification":"PASS","regression":"PASS",
        "ai_worker_model_invocation":"HOLD_NO_DSH_MODEL_PROVIDER_IN_GITHUB_CI",
        "runtime_and_toolchain":"REAL_EXECUTION_PASS",
    }
    report_paths=ReportRenderer().write_all(report,str(out/"report"))
    append_business_truth(truth,"REPORT_UPSERTED",{
        "id":"REPORT-RC7-REAL-DEMO","release_id":RELEASE,"mission_id":MISSION,"requirement_id":REQ,
        "status":"PASS","artifact_refs":[report_paths["json"],report_paths["html"]],"generated_at":int(time.time()),
    })
    summary={
        "status":"PASS","gate":"RC7_REAL_DEMO_RUNTIME_PASS","mission_id":MISSION,"release_id":RELEASE,
        "mission_state":mission_result.mission_state,"sufficiency_state":mission_result.sufficiency_state,
        "release_recommendation":mission_result.release_recommendation,"requirement":str(req_docx),
        "target_repo":str(repo),"base_commit":base_sha,"buggy_commit":buggy_sha,"fix_commit":fix_sha,
        "defect_id":defect.defect_id,"defect_state":defect.state.value,"r1":str(r1),"report":report_paths,
        "ai_worker_model_invocation":"HOLD_NO_DSH_MODEL_PROVIDER_IN_GITHUB_CI",
        "proof":"requirement->git/codegraph->test design->script readiness->real pytest fail->product defect->real reproduction->external fix->real verification/regression->defect close->mission converge->R1 report",
    }
    write_json(out/"demo-summary.json",summary)
    print(json.dumps(summary,ensure_ascii=False,sort_keys=True,indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
